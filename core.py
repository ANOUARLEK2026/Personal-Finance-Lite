"""Data layer for Personal Finance Lite (no GUI imports, so it is easy to test)."""
import calendar
import csv
import html
import json
import sqlite3
from datetime import date, datetime, timedelta
from pathlib import Path

APP_NAME = "Personal Finance Pro"
APP_VERSION = "2.0.0"
# Same data folder as Personal Finance Lite, so existing data carries over automatically.
APP_DIR = Path.home() / ".personal_finance_lite"
DB = APP_DIR / "finance.db"
BACKUP_DIR = APP_DIR / "backups"
MAX_BACKUPS = 15
AUTO_BACKUP_HOURS = 20

CATEGORIES = [
    "Food", "Transport", "Housing", "Bills", "Shopping",
    "Entertainment", "Health", "Education", "Salary", "Other"
]
KINDS = ("Expense", "Income")
CSV_HEADER = ["ID", "Type", "Amount", "Category", "Date", "Note"]
COLUMNS = "id,kind,amount,category,tx_date,note"


# --------------------------------------------------------------------------
# Database
# --------------------------------------------------------------------------
def db():
    APP_DIR.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    con.execute(
        """CREATE TABLE IF NOT EXISTS transactions(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kind TEXT NOT NULL,
            amount REAL NOT NULL,
            category TEXT,
            tx_date TEXT NOT NULL,
            note TEXT
        )"""
    )
    con.execute(
        """CREATE TABLE IF NOT EXISTS budgets(
            category TEXT PRIMARY KEY,
            amount REAL NOT NULL DEFAULT 0
        )"""
    )
    con.execute(
        """CREATE TABLE IF NOT EXISTS recurring(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            kind TEXT NOT NULL,
            amount REAL NOT NULL,
            category TEXT,
            note TEXT,
            day INTEGER NOT NULL DEFAULT 1,
            start_date TEXT NOT NULL,
            last_applied TEXT
        )"""
    )
    con.execute("CREATE INDEX IF NOT EXISTS idx_tx_date ON transactions(tx_date)")
    con.commit()
    return con


def money(value):
    return f"{value:,.2f} DZD"


def month_bounds(year, month):
    start = date(year, month, 1)
    if month == 12:
        end = date(year + 1, 1, 1) - timedelta(days=1)
    else:
        end = date(year, month + 1, 1) - timedelta(days=1)
    return start.isoformat(), end.isoformat()


def get_range(period, today=None):
    today = today or date.today()
    if period == "This month":
        return month_bounds(today.year, today.month)
    if period == "Last 30 days":
        return ((today - timedelta(days=29)).isoformat(), today.isoformat())
    if period == "This year":
        return (date(today.year, 1, 1).isoformat(), date(today.year, 12, 31).isoformat())
    return (None, None)


def fetch_transactions(start=None, end=None):
    con = db()
    query = f"SELECT {COLUMNS} FROM transactions WHERE 1=1"
    params = []
    if start:
        query += " AND tx_date >= ?"
        params.append(start)
    if end:
        query += " AND tx_date <= ?"
        params.append(end)
    query += " ORDER BY tx_date DESC, id DESC"
    rows = con.execute(query, params).fetchall()
    con.close()
    return rows


def get_transaction(ident):
    con = db()
    row = con.execute(f"SELECT {COLUMNS} FROM transactions WHERE id=?", (ident,)).fetchone()
    con.close()
    return row


def add_transaction(values):
    """Insert a transaction and return its new id."""
    con = db()
    cur = con.execute(
        "INSERT INTO transactions(kind,amount,category,tx_date,note) VALUES(?,?,?,?,?)",
        values,
    )
    con.commit()
    new_id = cur.lastrowid
    con.close()
    return new_id


def count_transactions():
    con = db()
    n = con.execute("SELECT COUNT(*) FROM transactions").fetchone()[0]
    con.close()
    return n


def update_transaction(ident, values):
    con = db()
    con.execute(
        "UPDATE transactions SET kind=?, amount=?, category=?, tx_date=?, note=? WHERE id=?",
        (*values, ident),
    )
    con.commit()
    con.close()


def delete_transactions(ids):
    ids = list(ids)
    if not ids:
        return
    con = db()
    con.executemany("DELETE FROM transactions WHERE id=?", [(i,) for i in ids])
    con.commit()
    con.close()


def clear_transactions():
    con = db()
    con.execute("DELETE FROM transactions")
    con.commit()
    con.close()


def load_budgets():
    con = db()
    data = {r["category"]: float(r["amount"]) for r in con.execute("SELECT category, amount FROM budgets")}
    con.close()
    return data


def save_budgets(values):
    con = db()
    for category, amount in values.items():
        con.execute(
            "INSERT INTO budgets(category, amount) VALUES(?, ?) "
            "ON CONFLICT(category) DO UPDATE SET amount=excluded.amount",
            (category, amount),
        )
    con.commit()
    con.close()


def extra_categories():
    """Categories typed by the user that are not in the default list."""
    con = db()
    rows = con.execute(
        "SELECT DISTINCT category FROM transactions WHERE category IS NOT NULL AND category != ''"
    ).fetchall()
    con.close()
    return sorted({r["category"] for r in rows} - set(CATEGORIES), key=str.lower)


# --------------------------------------------------------------------------
# Calculations
# --------------------------------------------------------------------------
def totals(rows):
    income = sum(float(r["amount"]) for r in rows if r["kind"] == "Income")
    expense = sum(float(r["amount"]) for r in rows if r["kind"] == "Expense")
    return income, expense


def category_spending(rows):
    spent = {}
    for r in rows:
        if r["kind"] == "Expense":
            cat = r["category"] or "Other"
            spent[cat] = spent.get(cat, 0.0) + float(r["amount"])
    return spent


def month_series(months=6, today=None):
    """[(label, income, expense)] for the last `months` months, oldest first."""
    today = today or date.today()
    keys = []
    y, m = today.year, today.month
    for _ in range(months):
        keys.append((y, m))
        m -= 1
        if m == 0:
            m, y = 12, y - 1
    keys.reverse()
    start = date(keys[0][0], keys[0][1], 1).isoformat()
    con = db()
    rows = con.execute(
        "SELECT substr(tx_date,1,7) AS ym, kind, SUM(amount) AS total "
        "FROM transactions WHERE tx_date >= ? GROUP BY ym, kind",
        (start,),
    ).fetchall()
    con.close()
    lookup = {(r["ym"], r["kind"]): float(r["total"]) for r in rows}
    out = []
    for y, m in keys:
        key = f"{y:04d}-{m:02d}"
        out.append((
            date(y, m, 1).strftime("%b"),
            lookup.get((key, "Income"), 0.0),
            lookup.get((key, "Expense"), 0.0),
        ))
    return out


def over_budget(budgets, today=None):
    """Categories whose spending this month exceeds their budget: [(cat, spent, budget)]."""
    today = today or date.today()
    start, end = month_bounds(today.year, today.month)
    spent = category_spending(fetch_transactions(start, end))
    result = [(c, spent.get(c, 0.0), b) for c, b in budgets.items() if b > 0 and spent.get(c, 0.0) > b]
    return sorted(result, key=lambda x: x[1] - x[2], reverse=True)


# --------------------------------------------------------------------------
# Backups
# --------------------------------------------------------------------------
def list_backups():
    if not BACKUP_DIR.exists():
        return []
    return sorted(BACKUP_DIR.glob("finance_*.db"), key=lambda p: p.stat().st_mtime, reverse=True)


def latest_backup():
    backups = list_backups()
    return backups[0] if backups else None


def _copy_db(src_path, dst_path):
    src = sqlite3.connect(src_path)
    dst = sqlite3.connect(dst_path)
    try:
        src.backup(dst)
    finally:
        dst.close()
        src.close()


def prune_backups():
    for old in list_backups()[MAX_BACKUPS:]:
        try:
            old.unlink()
        except OSError:
            pass


def make_backup(tag="manual"):
    """Create a consistent copy of the database. Returns the new path (or None)."""
    if not DB.exists():
        return None
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    dest = BACKUP_DIR / f"finance_{datetime.now():%Y%m%d_%H%M%S}_{tag}.db"
    _copy_db(DB, dest)
    prune_backups()
    return dest


def auto_backup_if_due():
    last = latest_backup()
    if last is not None:
        age = datetime.now() - datetime.fromtimestamp(last.stat().st_mtime)
        if age < timedelta(hours=AUTO_BACKUP_HOURS):
            return None
    return make_backup("auto")


def validate_backup(path):
    try:
        con = sqlite3.connect(Path(path).resolve().as_uri() + "?mode=ro", uri=True)
        try:
            con.execute("SELECT COUNT(*) FROM transactions").fetchone()
            con.execute("SELECT COUNT(*) FROM budgets").fetchone()
        finally:
            con.close()
        return True
    except sqlite3.Error:
        return False


def restore_backup(path):
    """Replace the current data with a backup. A safety copy of the current data is made first."""
    if not validate_backup(path):
        raise ValueError("This file is not a valid Personal Finance Lite backup.")
    make_backup("before-restore")
    _copy_db(path, DB)


# --------------------------------------------------------------------------
# CSV import / export
# --------------------------------------------------------------------------
def export_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(CSV_HEADER)
        for r in rows:
            writer.writerow([r["id"], r["kind"], r["amount"], r["category"] or "Other", r["tx_date"], r["note"] or ""])


def _parse_amount(text):
    s = str(text).strip().replace(" ", "").replace("DZD", "").replace("+", "")
    if "," in s and "." in s:
        s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    return float(s)


def _parse_date(text):
    s = str(text).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"bad date: {s!r}")


def import_csv(path):
    """Import transactions from a CSV (same columns as the export). Returns (added, duplicates, invalid)."""
    added = duplicates = invalid = 0
    con = db()
    existing = {
        (r["kind"], round(float(r["amount"]), 2), r["category"] or "Other", r["tx_date"], r["note"] or "")
        for r in con.execute("SELECT kind, amount, category, tx_date, note FROM transactions")
    }
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for raw in reader:
            row = {(k or "").strip().lower(): (v or "").strip() for k, v in raw.items()}
            try:
                kind = row.get("type", "").capitalize()
                if kind not in KINDS:
                    raise ValueError("bad type")
                amount = _parse_amount(row.get("amount", ""))
                if amount <= 0:
                    raise ValueError("bad amount")
                tx_date = _parse_date(row.get("date", ""))
            except ValueError:
                invalid += 1
                continue
            category = row.get("category") or "Other"
            note = row.get("note", "")
            key = (kind, round(amount, 2), category, tx_date, note)
            if key in existing:
                duplicates += 1
                continue
            existing.add(key)
            con.execute(
                "INSERT INTO transactions(kind,amount,category,tx_date,note) VALUES(?,?,?,?,?)",
                (kind, amount, category, tx_date, note),
            )
            added += 1
    con.commit()
    con.close()
    return added, duplicates, invalid


# --------------------------------------------------------------------------
# Undo support
# --------------------------------------------------------------------------
def restore_transactions(rows):
    """Re-insert previously deleted rows (keeps their original ids). Used by Undo."""
    con = db()
    con.executemany(
        "INSERT OR REPLACE INTO transactions(id,kind,amount,category,tx_date,note) VALUES(?,?,?,?,?,?)",
        [(r["id"], r["kind"], r["amount"], r["category"], r["tx_date"], r["note"]) for r in rows],
    )
    con.commit()
    con.close()



# --------------------------------------------------------------------------
# Insights: period comparison, averages, projection
# --------------------------------------------------------------------------
def _shift_month(d, delta):
    y, m = d.year, d.month + delta
    while m < 1:
        m, y = m + 12, y - 1
    while m > 12:
        m, y = m - 12, y + 1
    return y, m


def previous_range(period, today=None):
    """The comparable earlier window for a period name ('same elapsed days' for in-progress periods)."""
    today = today or date.today()
    if period == "This month":
        y, m = _shift_month(today, -1)
        last_day = calendar.monthrange(y, m)[1]
        return date(y, m, 1).isoformat(), date(y, m, min(today.day, last_day)).isoformat()
    if period == "Last 30 days":
        end = today - timedelta(days=30)
        return (end - timedelta(days=29)).isoformat(), end.isoformat()
    if period == "This year":
        y = today.year - 1
        last_day = calendar.monthrange(y, today.month)[1]
        return date(y, 1, 1).isoformat(), date(y, today.month, min(today.day, last_day)).isoformat()
    return (None, None)


def pct_change(current, previous):
    """Percent change, or None when there is nothing to compare against."""
    if previous is None or previous == 0:
        return None
    return (current - previous) / abs(previous) * 100


def daily_average_spend(rows, start=None, end=None):
    expense = totals(rows)[1]
    if not rows:
        return 0.0
    first = date.fromisoformat(start) if start else min(date.fromisoformat(r["tx_date"]) for r in rows)
    last = min(date.fromisoformat(end), date.today()) if end else date.today()
    days = max((last - first).days + 1, 1)
    return expense / days


def month_projection(today=None):
    """(spent so far this month, projected month-end spending)."""
    today = today or date.today()
    start, end = month_bounds(today.year, today.month)
    spent = totals(fetch_transactions(start, end))[1]
    days_in_month = calendar.monthrange(today.year, today.month)[1]
    return spent, spent / max(today.day, 1) * days_in_month


def top_expenses(rows, n=5):
    expenses = [r for r in rows if r["kind"] == "Expense"]
    return sorted(expenses, key=lambda r: float(r["amount"]), reverse=True)[:n]


def monthly_totals(rows):
    """[(YYYY-MM, income, expense)] for every month touched by `rows`, oldest first, gaps filled."""
    if not rows:
        return []
    agg = {}
    for r in rows:
        key = r["tx_date"][:7]
        inc, exp = agg.get(key, (0.0, 0.0))
        if r["kind"] == "Income":
            inc += float(r["amount"])
        else:
            exp += float(r["amount"])
        agg[key] = (inc, exp)
    first, last = min(agg), max(agg)
    y, m = int(first[:4]), int(first[5:7])
    out = []
    while f"{y:04d}-{m:02d}" <= last:
        key = f"{y:04d}-{m:02d}"
        out.append((key, *agg.get(key, (0.0, 0.0))))
        y, m = _shift_month(date(y, m, 1), 1)
    return out


# --------------------------------------------------------------------------
# Recurring transactions
# --------------------------------------------------------------------------
def list_recurring():
    con = db()
    rows = con.execute(
        "SELECT id,kind,amount,category,note,day,start_date,last_applied FROM recurring ORDER BY day, id"
    ).fetchall()
    con.close()
    return rows


def add_recurring(kind, amount, category, note, day, start_date):
    con = db()
    con.execute(
        "INSERT INTO recurring(kind,amount,category,note,day,start_date) VALUES(?,?,?,?,?,?)",
        (kind, amount, category, note, int(day), start_date),
    )
    con.commit()
    con.close()


def delete_recurring(ident):
    con = db()
    con.execute("DELETE FROM recurring WHERE id=?", (ident,))
    con.commit()
    con.close()


def apply_recurring(today=None):
    """Create the monthly transactions that are due. Safe to call on every start. Returns how many were added."""
    today = today or date.today()
    con = db()
    created = 0
    for rule in con.execute("SELECT * FROM recurring").fetchall():
        start = date.fromisoformat(rule["start_date"])
        y, m = start.year, start.month
        last_key = rule["last_applied"] or ""
        newest = last_key
        while (y, m) <= (today.year, today.month):
            key = f"{y:04d}-{m:02d}"
            if key > last_key:
                day = min(int(rule["day"]), calendar.monthrange(y, m)[1])
                when = date(y, m, day)
                if when >= start and when <= today:
                    con.execute(
                        "INSERT INTO transactions(kind,amount,category,tx_date,note) VALUES(?,?,?,?,?)",
                        (rule["kind"], rule["amount"], rule["category"], when.isoformat(),
                         (rule["note"] or "") or "Recurring"),
                    )
                    created += 1
                    newest = max(newest, key)
                elif when > today:
                    break
            y, m = _shift_month(date(y, m, 1), 1)
        if newest != last_key:
            con.execute("UPDATE recurring SET last_applied=? WHERE id=?", (newest, rule["id"]))
    con.commit()
    con.close()
    return created


# --------------------------------------------------------------------------
# Export: JSON
# --------------------------------------------------------------------------
def export_json(path, rows, budgets=None, filters_text=""):
    income, expense = totals(rows)
    payload = {
        "app": APP_NAME,
        "version": APP_VERSION,
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "filters": filters_text,
        "summary": {
            "income": round(income, 2),
            "expenses": round(expense, 2),
            "balance": round(income - expense, 2),
            "transactions": len(rows),
        },
        "budgets": budgets or {},
        "transactions": [
            {"id": r["id"], "type": r["kind"], "amount": float(r["amount"]),
             "category": r["category"] or "Other", "date": r["tx_date"], "note": r["note"] or ""}
            for r in rows
        ],
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


# --------------------------------------------------------------------------
# Export: Excel workbook (live formulas + native charts)
# --------------------------------------------------------------------------
def export_xlsx(path, rows, budgets=None, filters_text=""):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    FONT = "Arial"
    NAVY, SLATE, LIGHT = "0F172A", "475569", "E2E8F0"
    money_fmt = '#,##0.00;[Red]-#,##0.00;"-"'
    head_fill = PatternFill("solid", fgColor=NAVY)
    head_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    base = Font(name=FONT, size=10)
    bold = Font(name=FONT, size=10, bold=True)
    thin = Side(style="thin", color=LIGHT)
    box = Border(bottom=thin)

    def style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 22

    def set_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    ordered = sorted(rows, key=lambda r: (r["tx_date"], r["id"]))
    n = len(ordered)
    last = max(n + 1, 2)  # last data row on Transactions (formulas stay valid when empty)

    wb = Workbook()

    # ---------------- Transactions ------------------------------------
    tx = wb.active
    tx.title = "Transactions"
    tx.append(["Date", "Type", "Category", f"Amount ({'DZD'})", "Note", "ID"])
    style_header(tx, 1, 6)
    for r in ordered:
        tx.append([date.fromisoformat(r["tx_date"]), r["kind"], r["category"] or "Other",
                   float(r["amount"]), r["note"] or "", r["id"]])
    for row in tx.iter_rows(min_row=2, max_row=n + 1):
        for cell in row:
            cell.font, cell.border = base, box
        row[0].number_format = "yyyy-mm-dd"
        row[0].alignment = Alignment(horizontal="center")
        row[1].alignment = Alignment(horizontal="center")
        row[3].number_format = money_fmt
        row[1].font = Font(name=FONT, size=10, bold=True,
                           color="15803D" if row[1].value == "Income" else "BE123C")
    tx.freeze_panes = "A2"
    tx.auto_filter.ref = f"A1:F{max(n + 1, 2)}"
    set_widths(tx, [13, 11, 18, 18, 46, 8])

    rng = lambda col: f"Transactions!${col}$2:${col}${last}"

    # ---------------- Summary -----------------------------------------
    sm = wb.create_sheet("Summary", 0)
    sm.sheet_view.showGridLines = False
    sm["A1"] = f"{APP_NAME} — Report"
    sm["A1"].font = Font(name=FONT, size=18, bold=True, color=NAVY)
    sm["A2"] = f"Generated {datetime.now():%Y-%m-%d %H:%M}"
    sm["A3"] = f"Filters: {filters_text or 'none'}"
    for ref in ("A2", "A3"):
        sm[ref].font = Font(name=FONT, size=10, color=SLATE)
    kpis = [
        ("Income", f'=SUMIFS({rng("D")},{rng("B")},"Income")', money_fmt),
        ("Expenses", f'=SUMIFS({rng("D")},{rng("B")},"Expense")', money_fmt),
        ("Balance", "=B5-B6", money_fmt),
        ("Saving rate", "=IF(B5=0,0,B7/B5)", "0.0%"),
        ("Transactions", f'=COUNTIF({rng("B")},"Income")+COUNTIF({rng("B")},"Expense")', "#,##0"),
        ("Average expense", '=IF(COUNTIF(' + rng("B") + ',"Expense")=0,0,B6/COUNTIF(' + rng("B") + ',"Expense"))', money_fmt),
    ]
    sm["A4"], sm["B4"] = "Metric", "Value (DZD)"
    style_header(sm, 4, 2)
    for i, (label, formula, fmt) in enumerate(kpis, start=5):
        sm.cell(row=i, column=1, value=label).font = bold
        c = sm.cell(row=i, column=2, value=formula)
        c.number_format, c.font = fmt, base
        for col in (1, 2):
            sm.cell(row=i, column=col).border = box
    set_widths(sm, [26, 22, 4, 16, 16, 16])

    # Top expenses
    sm["A13"] = "Largest expenses"
    sm["A13"].font = Font(name=FONT, size=12, bold=True, color=NAVY)
    for c, h in enumerate(["Category / note", "Amount (DZD)", "", "Date"], 1):
        if h:
            sm.cell(row=14, column=c, value=h)
    for c in (1, 2, 4):
        cell = sm.cell(row=14, column=c)
        cell.fill, cell.font = head_fill, head_font
    for i, r in enumerate(top_expenses(rows, 5), start=15):
        label = (r["category"] or "Other") + (f" — {r['note']}" if r["note"] else "")
        sm.cell(row=i, column=1, value=label).font = base
        c = sm.cell(row=i, column=2, value=float(r["amount"]))
        c.number_format, c.font = money_fmt, base
        d = sm.cell(row=i, column=4, value=date.fromisoformat(r["tx_date"]))
        d.number_format, d.font = "yyyy-mm-dd", base

    # ---------------- By Category -------------------------------------
    spent = sorted(category_spending(rows).items(), key=lambda x: x[1], reverse=True)
    cat = wb.create_sheet("By Category")
    cat.append(["Category", "Spent (DZD)", "Share", "Expenses (count)"])
    style_header(cat, 1, 4)
    for i, (name, _) in enumerate(spent, start=2):
        cat.cell(row=i, column=1, value=name)
        cat.cell(row=i, column=2,
                 value=f'=SUMIFS({rng("D")},{rng("B")},"Expense",{rng("C")},A{i})')
        cat.cell(row=i, column=3, value=f"=IF(Summary!$B$6=0,0,B{i}/Summary!$B$6)")
        cat.cell(row=i, column=4, value=f'=COUNTIFS({rng("B")},"Expense",{rng("C")},A{i})')
        for col in range(1, 5):
            cat.cell(row=i, column=col).font = base
            cat.cell(row=i, column=col).border = box
        cat.cell(row=i, column=2).number_format = money_fmt
        cat.cell(row=i, column=3).number_format = "0.0%"
    set_widths(cat, [22, 18, 10, 17])
    if spent:
        pie = PieChart()
        pie.title = "Spending by category"
        pie.add_data(Reference(cat, min_col=2, min_row=1, max_row=len(spent) + 1), titles_from_data=True)
        pie.set_categories(Reference(cat, min_col=1, min_row=2, max_row=len(spent) + 1))
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.height, pie.width = 9, 13
        cat.add_chart(pie, "F2")

    # ---------------- Monthly -----------------------------------------
    mo = wb.create_sheet("Monthly")
    mo.append(["Month", "Income (DZD)", "Expenses (DZD)", "Net (DZD)", "Saving rate"])
    style_header(mo, 1, 5)
    months = monthly_totals(rows)
    for i, (key, _, _) in enumerate(months, start=2):
        mo.cell(row=i, column=1, value=date(int(key[:4]), int(key[5:7]), 1)).number_format = "mmm yyyy"
        mo.cell(row=i, column=2,
                value=f'=SUMIFS({rng("D")},{rng("B")},"Income",{rng("A")},">="&A{i},{rng("A")},"<"&EDATE(A{i},1))')
        mo.cell(row=i, column=3,
                value=f'=SUMIFS({rng("D")},{rng("B")},"Expense",{rng("A")},">="&A{i},{rng("A")},"<"&EDATE(A{i},1))')
        mo.cell(row=i, column=4, value=f"=B{i}-C{i}")
        mo.cell(row=i, column=5, value=f"=IF(B{i}=0,0,D{i}/B{i})")
        for col in range(1, 6):
            mo.cell(row=i, column=col).font = base
            mo.cell(row=i, column=col).border = box
        for col in (2, 3, 4):
            mo.cell(row=i, column=col).number_format = money_fmt
        mo.cell(row=i, column=5).number_format = "0.0%"
    set_widths(mo, [14, 18, 18, 18, 13])
    if months:
        bar = BarChart()
        bar.type, bar.title = "col", "Income vs expenses"
        bar.add_data(Reference(mo, min_col=2, max_col=3, min_row=1, max_row=len(months) + 1), titles_from_data=True)
        bar.set_categories(Reference(mo, min_col=1, min_row=2, max_row=len(months) + 1))
        bar.height, bar.width = 9, 16
        mo.add_chart(bar, "G2")

    wb.calculation.fullCalcOnLoad = True  # Excel recalculates the live formulas when the file opens
    wb.save(path)


# --------------------------------------------------------------------------
# Printable report (PDF) — HTML for QTextDocument
# --------------------------------------------------------------------------
def build_report_html(rows, filters_text, bar_img=None, donut_img=None):
    income, expense = totals(rows)
    esc = html.escape
    spent = sorted(category_spending(rows).items(), key=lambda x: x[1], reverse=True)
    rate = ((income - expense) / income * 100) if income else 0
    th = "style='background-color:#0f172a;color:#ffffff;'"
    zebra = ("#f8fafc", "#ffffff")

    def kpi(title, value, color):
        return (f"<td width='25%' style='background-color:#f1f5f9;border-left:4px solid {color};'>"
                f"<span style='color:#64748b;font-size:8pt;'>{title}</span><br>"
                f"<b style='font-size:12pt;color:#0f172a;'>{value}</b></td>")

    cat_rows = "".join(
        f"<tr bgcolor='{zebra[i % 2]}'><td>{esc(c)}</td><td align='right'>{money(v)}</td>"
        f"<td align='right'>{(v / expense * 100 if expense else 0):.1f}%</td></tr>"
        for i, (c, v) in enumerate(spent)
    ) or "<tr><td colspan='3'>No expenses in this selection.</td></tr>"

    tx_rows = "".join(
        f"<tr bgcolor='{zebra[i % 2]}'><td>{esc(r['tx_date'])}</td><td>{esc(r['kind'])}</td>"
        f"<td>{esc(r['category'] or 'Other')}</td>"
        f"<td align='right' style='color:{'#15803d' if r['kind'] == 'Income' else '#be123c'};'>"
        f"{'+' if r['kind'] == 'Income' else '-'}{money(float(r['amount']))}</td>"
        f"<td>{esc(r['note'] or '')}</td></tr>"
        for i, r in enumerate(rows)
    ) or "<tr><td colspan='5'>No transactions.</td></tr>"

    charts = ""
    if bar_img or donut_img:
        cells = ""
        if bar_img:
            cells += f"<td width='58%'><img src='{bar_img}' width='380'></td>"
        if donut_img:
            cells += f"<td width='42%'><img src='{donut_img}' width='270'></td>"
        charts = f"<table width='100%' cellspacing='0' cellpadding='4'><tr>{cells}</tr></table>"

    top = "".join(
        f"<tr bgcolor='{zebra[i % 2]}'><td>{esc(r['tx_date'])}</td><td>{esc(r['category'] or 'Other')}</td>"
        f"<td>{esc(r['note'] or '')}</td><td align='right'>{money(float(r['amount']))}</td></tr>"
        for i, r in enumerate(top_expenses(rows, 5))
    )
    top_block = (
        f"<h3 style='color:#0f172a;'>Largest expenses</h3>"
        f"<table width='100%' cellspacing='0' cellpadding='5'>"
        f"<tr><th {th} align='left'>Date</th><th {th} align='left'>Category</th><th {th} align='left'>Note</th>"
        f"<th {th} align='right'>Amount</th></tr>{top}</table>"
    ) if top else ""

    return f"""
    <h1 style='margin-bottom:0;color:#0f172a;'>{APP_NAME} &mdash; Financial Report</h1>
    <p style='color:#64748b;'>Generated {datetime.now():%Y-%m-%d %H:%M} &nbsp;|&nbsp; Filters: {esc(filters_text)}</p>
    <table width='100%' cellspacing='8' cellpadding='9'><tr>
      {kpi('INCOME', money(income), '#22c55e')}
      {kpi('EXPENSES', money(expense), '#fb7185')}
      {kpi('BALANCE', money(income - expense), '#60a5fa')}
      {kpi('SAVING RATE', f'{rate:.1f}%', '#a78bfa')}
    </tr></table>
    {charts}
    <h3 style='color:#0f172a;'>Spending by category</h3>
    <table width='100%' cellspacing='0' cellpadding='5'>
      <tr><th {th} align='left'>Category</th><th {th} align='right'>Spent</th><th {th} align='right'>Share</th></tr>
      {cat_rows}
    </table>
    {top_block}
    <h3 style='color:#0f172a;'>All transactions ({len(rows)})</h3>
    <table width='100%' cellspacing='0' cellpadding='4'>
      <thead><tr><th {th} align='left'>Date</th><th {th} align='left'>Type</th><th {th} align='left'>Category</th>
          <th {th} align='right'>Amount</th><th {th} align='left'>Note</th></tr></thead>
      {tx_rows}
    </table>
    """
